import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    base_dir = os.path.dirname(os.path.abspath(__file__)) if '__file__' in globals() else os.getcwd()
    assets_dir = os.path.join(base_dir, "docs", "assets")
    if not os.path.exists(assets_dir):
        assets_dir = os.path.join(base_dir, "assets")

    img_hero = os.path.join(assets_dir, "website_hero_landing.jpg")
    img_swagger = os.path.join(assets_dir, "website_swagger_docs.jpg")
    img_datasets = os.path.join(assets_dir, "website_datasets_explorer.jpg")
    img_dashboard = os.path.join(assets_dir, "website_dashboard_profile.jpg")

    # --- COLOR PALETTE (Cosmic Dark / Modern Enterprise) ---
    BG_DARK = RGBColor(10, 15, 29)          # #0A0F1D Deep Cosmic Navy
    CARD_BG = RGBColor(19, 29, 49)          # #131D31 Glass slate card
    CARD_BORDER = RGBColor(40, 56, 84)      # #283854 Subtle border
    CARD_HIGHLIGHT = RGBColor(28, 41, 68)   # #1C2944
    
    CYAN = RGBColor(56, 189, 248)           # #38BDF8 Electric Cyan
    PURPLE = RGBColor(168, 85, 247)         # #A855F7 Cosmic Purple
    GOLD = RGBColor(245, 158, 11)           # #F59E0B Solar Amber
    EMERALD = RGBColor(16, 185, 129)        # #10B981 Emerald
    ROSE = RGBColor(244, 63, 94)            # #F43F5E Neon Rose
    
    TEXT_WHITE = RGBColor(248, 250, 252)    # #F8FAFC Pure bright
    TEXT_MUTED = RGBColor(203, 213, 225)    # #CBD5E1 Slate light
    TEXT_DIM = RGBColor(148, 163, 184)      # #94A3B8 Dim slate

    def apply_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_DARK
        bg.line.fill.background()
        
        # Subtle header accent strip at top
        strip = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.08))
        strip.fill.solid()
        strip.fill.fore_color.rgb = CYAN
        strip.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="MYSTICAI CAPSTONE PROJECT", tag_color=CYAN, slide_num=None):
        # Category Tag
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(9.0), Inches(0.28))
        tf_tag = tag_box.text_frame
        tf_tag.word_wrap = True
        tf_tag.margin_left = tf_tag.margin_right = tf_tag.margin_top = tf_tag.margin_bottom = 0
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = category_text.upper()
        p_tag.font.size = Pt(10.5)
        p_tag.font.bold = True
        p_tag.font.color.rgb = tag_color

        # Slide Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.65), Inches(9.5), Inches(0.6))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = tf_title.margin_right = tf_title.margin_top = tf_title.margin_bottom = 0
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(21)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_WHITE

        # Slide Number Badge
        if slide_num:
            num_box = slide.shapes.add_textbox(Inches(11.3), Inches(0.45), Inches(1.2), Inches(0.4))
            tf_num = num_box.text_frame
            p_num = tf_num.paragraphs[0]
            p_num.text = f"0{slide_num} / 07"
            p_num.alignment = PP_ALIGN.RIGHT
            p_num.font.size = Pt(13)
            p_num.font.bold = True
            p_num.font.color.rgb = TEXT_DIM

        # Decorative Divider Line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.3), Inches(11.733), Inches(0.02))
        line.fill.solid()
        line.fill.fore_color.rgb = CARD_BORDER
        line.line.fill.background()

    def add_screenshot_frame(slide, img_path, left, top, width, height, title="LIVE SYSTEM SCREENSHOT", badge="VERIFIED UI"):
        # Outer Card
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left - 0.08), Inches(top - 0.08), Inches(width + 0.16), Inches(height + 0.55)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1.5)

        # Title Label Bar
        tb = slide.shapes.add_textbox(Inches(left), Inches(top - 0.02), Inches(width * 0.7), Inches(0.3))
        tf = tb.text_frame
        tf.margin_left = tf.margin_top = 0
        p = tf.paragraphs[0]
        p.text = f"🖥️  {title}"
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = CYAN

        # Badge
        badge_box = slide.shapes.add_textbox(Inches(left + width - 1.6), Inches(top - 0.02), Inches(1.6), Inches(0.3))
        tf_b = badge_box.text_frame
        tf_b.margin_right = tf_b.margin_top = 0
        pb = tf_b.paragraphs[0]
        pb.text = f"● {badge}"
        pb.alignment = PP_ALIGN.RIGHT
        pb.font.size = Pt(9)
        pb.font.bold = True
        pb.font.color.rgb = EMERALD

        # Image insertion
        if os.path.exists(img_path):
            slide.shapes.add_picture(img_path, Inches(left), Inches(top + 0.3), Inches(width), Inches(height - 0.05))

    # =========================================================================
    # SLIDE 1: Vision, Problem & Executive Overview
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    apply_bg(slide1)
    add_header(slide1, "MysticAI: Next-Gen Palmistry & Tarot Intelligence", "INFOSYS SPRINGBOARD CAPSTONE", CYAN, slide_num=1)

    # Left Column: Overview & Value Prop Cards
    # Hero Title / Tagline Card
    hero_card = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.8), Inches(1.65))
    hero_card.fill.solid()
    hero_card.fill.fore_color.rgb = CARD_BG
    hero_card.line.color.rgb = CARD_BORDER
    hero_card.line.width = Pt(1.5)

    tb_hero = slide1.shapes.add_textbox(Inches(1.0), Inches(1.62), Inches(5.4), Inches(1.4))
    tf_hero = tb_hero.text_frame
    tf_hero.word_wrap = True
    p1 = tf_hero.paragraphs[0]
    p1.text = "Multimodal AI Spiritual Intelligence Platform"
    p1.font.size = Pt(15)
    p1.font.bold = True
    p1.font.color.rgb = GOLD

    p2 = tf_hero.add_paragraph()
    p2.text = "Bridging ancient esoteric divination (Tarot & Palmistry) with modern Computer Vision, Google Gemini 2.0 Flash LLM, and full-stack reactive microservices."
    p2.font.size = Pt(11)
    p2.font.color.rgb = TEXT_MUTED

    # 4 Key Innovation Cards on the left
    features = [
        ("🃏 3D Interactive Card Draw", "Physics-driven visual shuffle, fanning & real-time 3D flip card picks", CYAN),
        ("🤖 Gemini Conversational AI", "Context-aware spiritual chatbot with domain-tuned guidance prompts", PURPLE),
        ("✋ CV Palm Line Intelligence", "Automated line detection, feature extraction & Big Five personality mapping", EMERALD),
        ("🛡️ Enterprise Polyglot Core", "FastAPI async engine + Next.js 14 + PostgreSQL, MongoDB & Redis", ROSE)
    ]

    card_y = 3.3
    for title, desc, col in features:
        box = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(card_y), Inches(5.8), Inches(0.88))
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = CARD_BORDER
        box.line.width = Pt(1)

        tb = slide1.shapes.add_textbox(Inches(0.98), Inches(card_y + 0.08), Inches(5.45), Inches(0.72))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = 0

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(11.5)
        p.font.bold = True
        p.font.color.rgb = col

        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.size = Pt(10)
        p_desc.font.color.rgb = TEXT_MUTED

        card_y += 0.98

    # Right Column: Hero Screenshot & Stat Chips
    add_screenshot_frame(slide1, img_hero, 6.9, 1.5, 5.6, 4.3, "LIVE PLATFORM HERO & INTERACTION PORTAL", "PRODUCTION READY")

    # Bottom Stat Row on Right
    stat_box1 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(6.22), Inches(2.68), Inches(0.95))
    stat_box1.fill.solid()
    stat_box1.fill.fore_color.rgb = CARD_HIGHLIGHT
    stat_box1.line.color.rgb = CYAN
    stat_box1.line.width = Pt(1)

    tb_s1 = slide1.shapes.add_textbox(Inches(7.0), Inches(6.28), Inches(2.48), Inches(0.8))
    tf_s1 = tb_s1.text_frame
    tf_s1.word_wrap = True
    p = tf_s1.paragraphs[0]
    p.text = "78 Cards · 6 Spreads"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = CYAN
    p_sub = tf_s1.add_paragraph()
    p_sub.text = "Full Major & Minor Arcana graph"
    p_sub.font.size = Pt(9.5)
    p_sub.font.color.rgb = TEXT_MUTED

    stat_box2 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.82), Inches(6.22), Inches(2.68), Inches(0.95))
    stat_box2.fill.solid()
    stat_box2.fill.fore_color.rgb = CARD_HIGHLIGHT
    stat_box2.line.color.rgb = PURPLE
    stat_box2.line.width = Pt(1)

    tb_s2 = slide1.shapes.add_textbox(Inches(9.92), Inches(6.28), Inches(2.48), Inches(0.8))
    tf_s2 = tb_s2.text_frame
    tf_s2.word_wrap = True
    p = tf_s2.paragraphs[0]
    p.text = "4-Tier RBAC & JWT"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = PURPLE
    p_sub = tf_s2.add_paragraph()
    p_sub.text = "Enterprise security & audit logging"
    p_sub.font.size = Pt(9.5)
    p_sub.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 2: Polyglot Architecture & Microservices Engine
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    apply_bg(slide2)
    add_header(slide2, "System Architecture: Dual-Database Polyglot Engine", "SYSTEM DESIGN & DATA FLOW", PURPLE, slide_num=2)

    # 3 Architecture Tier Cards on Left
    tiers = [
        ("⚡ Frontend Client Layer (Next.js 14)", [
            "React 18 + Next.js App Router for dynamic routing & SEO",
            "TailwindCSS + Framer Motion for 3D card flips & glassmorphism",
            "Real-time state hydration with JWT bearer interceptors"
        ], CYAN),
        ("🚀 Async Application Layer (FastAPI)", [
            "High-throughput Python 3.11 asynchronous REST microservices",
            "Strict Pydantic v2 contract validation for all payloads",
            "Modular routers: Auth, Readings, AI Chat, Datasets, Analytics"
        ], PURPLE),
        ("🗄️ Polyglot Dual-Database Layer", [
            "PostgreSQL (SQLAlchemy): ACID user auth, roles & transactions",
            "MongoDB (Motor): Flexible NoSQL reading logs, card spread graphs",
            "Redis: Token blacklisting, session caching & rate limiting"
        ], EMERALD)
    ]

    ty = 1.5
    for title, points, col in tiers:
        card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(ty), Inches(5.8), Inches(1.7))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1)

        tb = slide2.shapes.add_textbox(Inches(0.98), Inches(ty + 0.1), Inches(5.45), Inches(1.5))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = 0

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(12)
        p_t.font.bold = True
        p_t.font.color.rgb = col

        for pt in points:
            p_pt = tf.add_paragraph()
            p_pt.text = f"▪  {pt}"
            p_pt.font.size = Pt(9.5)
            p_pt.font.color.rgb = TEXT_MUTED

        ty += 1.85

    # Right: Swagger UI Screenshot and Containerization card
    add_screenshot_frame(slide2, img_swagger, 6.9, 1.5, 5.6, 3.8, "INTERACTIVE FASTAPI SWAGGER / REDOC API DOCS", "FULLY DOCUMENTED")

    # Bottom Right: Docker Compose & Infrastructure Summary Card
    infra_card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(5.7), Inches(5.6), Inches(1.45))
    infra_card.fill.solid()
    infra_card.fill.fore_color.rgb = CARD_HIGHLIGHT
    infra_card.line.color.rgb = CARD_BORDER
    infra_card.line.width = Pt(1)

    tb_inf = slide2.shapes.add_textbox(Inches(7.1), Inches(5.8), Inches(5.2), Inches(1.25))
    tf_inf = tb_inf.text_frame
    tf_inf.word_wrap = True
    p = tf_inf.paragraphs[0]
    p.text = "🐳 Dockerized Microservices & Unified Orchestration"
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = GOLD

    p1 = tf_inf.add_paragraph()
    p1.text = "▪ Single-command initialization via `docker-compose up -d`"
    p1.font.size = Pt(9.5)
    p1.font.color.rgb = TEXT_MUTED

    p2 = tf_inf.add_paragraph()
    p2.text = "▪ Isolated container networks for PostgreSQL (5432), MongoDB (27017), Redis (6379) & FastAPI (8000)"
    p2.font.size = Pt(9.5)
    p2.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 3: Core AI Engines — 3D Interactive Tarot & Palmistry
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    apply_bg(slide3)
    add_header(slide3, "Core AI Engines: 3D Interactive Tarot & Palm CV", "MULTIMODAL INTELLIGENCE", GOLD, slide_num=3)

    # Left: Datasets Screenshot
    add_screenshot_frame(slide3, img_datasets, 0.8, 1.5, 5.7, 4.0, "DIGITIZED KNOWLEDGE BASE & SPREAD EXPLORER", "78 CARDS + PALM DATA")

    # Bottom Left: Spread Systems Box
    spread_box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.9), Inches(5.7), Inches(1.25))
    spread_box.fill.solid()
    spread_box.fill.fore_color.rgb = CARD_HIGHLIGHT
    spread_box.line.color.rgb = GOLD
    spread_box.line.width = Pt(1)

    tb_sp = slide3.shapes.add_textbox(Inches(0.95), Inches(5.98), Inches(5.4), Inches(1.1))
    tf_sp = tb_sp.text_frame
    tf_sp.word_wrap = True
    p = tf_sp.paragraphs[0]
    p.text = "🔮 6 Standard Spread Frameworks Supported:"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = GOLD

    p_sp = tf_sp.add_paragraph()
    p_sp.text = "Single Card · Past-Present-Future · Celtic Cross (10 Cards) · Career Path · Love & Relationship · Horseshoe"
    p_sp.font.size = Pt(9.5)
    p_sp.font.color.rgb = TEXT_MUTED

    # Right: 2 Deep-Dive Cards (Tarot Engine + Palmistry CV Engine)
    # Tarot Engine Card
    t_card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.5), Inches(5.7), Inches(2.6))
    t_card.fill.solid()
    t_card.fill.fore_color.rgb = CARD_BG
    t_card.line.color.rgb = CARD_BORDER
    t_card.line.width = Pt(1)

    tb_t = slide3.shapes.add_textbox(Inches(7.0), Inches(1.6), Inches(5.3), Inches(2.4))
    tf_t = tb_t.text_frame
    tf_t.word_wrap = True
    p = tf_t.paragraphs[0]
    p.text = "🃏 3D Physics Interactive Tarot Engine"
    p.font.size = Pt(12.5)
    p.font.bold = True
    p.font.color.rgb = CYAN

    tarot_points = [
        "Interactive Card Draw: Physics-based visual deck shuffle, card fanning, and manual one-by-one pick flow",
        "3D CSS & Framer Flip: Recreates authentic physical tactile experience with flip animation & upright/reversed state",
        "Elemental & Astrological Mapping: Cards mapped to Fire, Water, Air, Earth, Zodiac signs, and planetary rulers",
        "Deterministic + AI Synthesis: Position-specific meanings fused into dynamic multi-card narratives"
    ]
    for pt in tarot_points:
        p_pt = tf_t.add_paragraph()
        p_pt.text = f"▪  {pt}"
        p_pt.font.size = Pt(9.5)
        p_pt.font.color.rgb = TEXT_MUTED

    # Palmistry CV Card
    p_card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(4.3), Inches(5.7), Inches(2.85))
    p_card.fill.solid()
    p_card.fill.fore_color.rgb = CARD_BG
    p_card.line.color.rgb = CARD_BORDER
    p_card.line.width = Pt(1)

    tb_p = slide3.shapes.add_textbox(Inches(7.0), Inches(4.4), Inches(5.3), Inches(2.6))
    tf_p = tb_p.text_frame
    tf_p.word_wrap = True
    p = tf_p.paragraphs[0]
    p.text = "✋ Computer Vision Palm Line Detection"
    p.font.size = Pt(12.5)
    p.font.bold = True
    p.font.color.rgb = EMERALD

    palm_points = [
        "5 Major/Minor Line Analysis: Heart (Emotions), Head (Intellect), Life (Vitality), Fate (Destiny), Sun (Creativity)",
        "Morphological Feature Extraction: Quantifies length, curvature, depth, fork branches, and line breaks",
        "Personality Archetype Engine: Mapped into Big Five personality metrics (Openness, Conscientiousness, Extroversion, etc.)",
        "Dual Input Mode: Supports real-time camera image upload or manual interactive parameter calibration"
    ]
    for pt in palm_points:
        p_pt = tf_p.add_paragraph()
        p_pt.text = f"▪  {pt}"
        p_pt.font.size = Pt(9.5)
        p_pt.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 4: Gemini 2.0 Conversational AI & Multimodal Synthesis
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    apply_bg(slide4)
    add_header(slide4, "Multimodal AI & Gemini 2.0 Conversational Chatbot", "COGNITIVE ARCHITECTURE", ROSE, slide_num=4)

    # 4 Architecture Pillars in a 2x2 Grid
    pillars = [
        ("🤖 Google Gemini 2.0 Flash Integration", [
            "Low-latency streaming inferences via `@google/genai` API",
            "System instruction engineering enforcing empathetic, mystical yet grounded tone",
            "Structured JSON response enforcement for cards, traits & recommendations"
        ], CYAN, 0.8, 1.5),
        ("🧠 Context-Aware Memory & Persona", [
            "Multi-turn conversation history retained with session sliding window",
            "Dynamic user profile injection (Zodiac sign, life path, previous readings)",
            "Dual-mode execution: Live Gemini LLM or Offline Rule-based Template Engine"
        ], PURPLE, 6.8, 1.5),
        ("🔮 Deep Narrative Synthesis Engine", [
            "Cross-modal fusion: Synthesizes Tarot card archetypes + Palmistry lines",
            "Multi-dimensional scoring across 4 life pillars: Career, Love, Health & Spirit",
            "Dynamic confidence scoring based on card harmony & line strength"
        ], GOLD, 0.8, 4.3),
        ("💡 Personalized Recommendation Engine", [
            "Actionable daily affirmations & mindfulness spiritual practices",
            "Crystal & gemstone pairings with elemental balancing advice",
            "Automated PDF & Word document generation for downloadable personal dossiers"
        ], EMERALD, 6.8, 4.3)
    ]

    for title, points, col, x, y in pillars:
        card = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(5.7), Inches(2.65))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1)

        tb = slide4.shapes.add_textbox(Inches(x + 0.2), Inches(y + 0.12), Inches(5.3), Inches(2.45))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = 0

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = col

        for pt in points:
            p_pt = tf.add_paragraph()
            p_pt.text = f"▪  {pt}"
            p_pt.font.size = Pt(9.5)
            p_pt.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 5: Enterprise RBAC, Analytics & Life Trend Intelligence
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    apply_bg(slide5)
    add_header(slide5, "Enterprise Security, 4-Tier RBAC & Analytics Dashboard", "SECURITY & ANALYTICS", EMERALD, slide_num=5)

    # Left: Dashboard Screenshot
    add_screenshot_frame(slide5, img_dashboard, 0.8, 1.5, 5.7, 4.0, "USER ANALYTICS & PROFILE DASHBOARD", "4 ROLES ACTIVE")

    # Bottom Left: Security Specs
    sec_box = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.9), Inches(5.7), Inches(1.25))
    sec_box.fill.solid()
    sec_box.fill.fore_color.rgb = CARD_HIGHLIGHT
    sec_box.line.color.rgb = EMERALD
    sec_box.line.width = Pt(1)

    tb_sec = slide5.shapes.add_textbox(Inches(0.95), Inches(5.98), Inches(5.4), Inches(1.1))
    tf_sec = tb_sec.text_frame
    tf_sec.word_wrap = True
    p = tf_sec.paragraphs[0]
    p.text = "🔐 Enterprise Security & Token Lifecycle:"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = EMERALD

    p_sec = tf_sec.add_paragraph()
    p_sec.text = "Bcrypt password hashing · HMAC-SHA256 JWT tokens · Redis token blacklisting on logout · CORS & Rate limiters"
    p_sec.font.size = Pt(9.5)
    p_sec.font.color.rgb = TEXT_MUTED

    # Right: 4-Tier RBAC Matrix & Analytics Intelligence
    # RBAC Table Card
    rbac_card = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.5), Inches(5.7), Inches(2.7))
    rbac_card.fill.solid()
    rbac_card.fill.fore_color.rgb = CARD_BG
    rbac_card.line.color.rgb = CARD_BORDER
    rbac_card.line.width = Pt(1)

    tb_r = slide5.shapes.add_textbox(Inches(7.0), Inches(1.6), Inches(5.3), Inches(2.5))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True
    p = tf_r.paragraphs[0]
    p.text = "👥 4-Tier Role-Based Access Control (RBAC)"
    p.font.size = Pt(12.5)
    p.font.bold = True
    p.font.color.rgb = CYAN

    roles = [
        ("🔮 Seeker (End User)", "Perform personal Tarot & Palm readings, chat with AI, view history"),
        ("🃏 Tarot Reader", "Manage client readings, customize interpretations, annotate spreads"),
        ("🌿 Consultant", "Access holistic life trend reports, suggest lifestyle remedies & crystals"),
        ("👑 Administrator", "Full user CRUD, system metrics, dataset moderation & audit logs")
    ]
    for r_title, r_desc in roles:
        p_r = tf_r.add_paragraph()
        p_r.text = f"{r_title}: "
        p_r.font.size = Pt(9.5)
        p_r.font.bold = True
        p_r.font.color.rgb = TEXT_WHITE
        # Add normal text to same paragraph
        run = p_r.add_run()
        run.text = r_desc
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED

    # Analytics Card
    an_card = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(4.4), Inches(5.7), Inches(2.75))
    an_card.fill.solid()
    an_card.fill.fore_color.rgb = CARD_BG
    an_card.line.color.rgb = CARD_BORDER
    an_card.line.width = Pt(1)

    tb_an = slide5.shapes.add_textbox(Inches(7.0), Inches(4.5), Inches(5.3), Inches(2.5))
    tf_an = tb_an.text_frame
    tf_an.word_wrap = True
    p = tf_an.paragraphs[0]
    p.text = "📈 Life Trend Analytics & Opportunity Forecasting"
    p.font.size = Pt(12.5)
    p.font.bold = True
    p.font.color.rgb = PURPLE

    analytics_points = [
        "Historical Pattern Detection: Tracks recurring Arcana suits (Cups, Wands, Swords, Pentacles) across user timeline",
        "Emotional & Career Forecasts: Identifies upcoming astrological peaks & growth opportunities",
        "Confidence & Harmony Scoring: Real-time statistical confidence metrics based on line clarity & card synergy",
        "Multi-Format Reporting: Instant one-click export to PDF / DOCX full spiritual dossiers"
    ]
    for pt in analytics_points:
        p_pt = tf_an.add_paragraph()
        p_pt.text = f"▪  {pt}"
        p_pt.font.size = Pt(9.5)
        p_pt.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 6: Complete Milestones Delivery Matrix (Milestones 1 - 4)
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    apply_bg(slide6)
    add_header(slide6, "Comprehensive Milestone Delivery: 100% Completion (M1 - M4)", "INFOSYS SPRINGBOARD MILESTONES REVIEW", EMERALD, slide_num=6)

    milestones = [
        ("MILESTONE 1: Foundations, Polyglot DB & RBAC", CYAN, 0.8, 1.5, [
            ("⚡ FastAPI Async Core", "30+ production REST endpoints with strict Pydantic v2 schemas"),
            ("🗄️ Dual-Database Layer", "PostgreSQL (ACID) + MongoDB (NoSQL) + Redis (Cache/Blacklist) with SQLite fallback"),
            ("🔐 Enterprise Security", "Bcrypt hashing, HMAC-SHA256 JWT tokens, Redis token blacklist & CORS middleware"),
            ("👥 4-Tier RBAC Engine", "Seeker, Tarot Reader, Consultant & Super Admin with 12 granular permissions"),
            ("📑 OpenAPI 3.0 Contracts", "Interactive Swagger UI & ReDoc specs with 100% endpoint documentation parity")
        ]),
        ("MILESTONE 2: Multimodal Engines & 1,000-Image Catalog", PURPLE, 6.78, 1.5, [
            ("✋ CV Palm Analysis Engine", "5 lines (Heart, Head, Life, Fate, Sun) with curvature, clarity & depth extraction"),
            ("🌿 Hand Morphology Analysis", "Earth, Air, Fire, Water elemental hand classifications & trait mapping"),
            ("🃏 78-Card Tarot Deck & Spreads", "6 distinct spread architectures, position-specific semantics & reverse orientations"),
            ("📂 1,000-Image Dataset Catalog", "500 labeled palms + 500 multi-style tarot cards with search & filter APIs"),
            ("📷 Dual Input Mode", "Live camera scanner with HUD overlay guide + multipart image file upload")
        ]),
        ("MILESTONE 3: AI Interpretation, Personality & Trends", GOLD, 0.8, 4.35, [
            ("🤖 Gemini 2.0 Flash LLM", "Domain-tuned prompt engineering for empathetic, mystical & grounded insights"),
            ("🔮 Cross-Modal Synthesis", "Deep narrative fusion synthesizing biometric palm lines with tarot card archetypes"),
            ("🧠 Big Five Personality Engine", "OCEAN trait modeling + 10 spiritual archetypes with strengths & behavioral patterns"),
            ("📈 Life Trend Forecasting", "Tracks 5 life dimensions (Emotional, Career, Spiritual, Vitality, Creative)"),
            ("⚖️ 5-Factor Weighted Scoring", "Normalized 0-100 score (30% Palm, 25% Tarot, 20% Persona, 15% Context, 10% Consistency)")
        ]),
        ("MILESTONE 4: Interactive UX, Dashboards & Quality", EMERALD, 6.78, 4.35, [
            ("🃏 3D Physics Card Draw", "Interactive visual shuffle, fanning, one-by-one manual pick & 3D CSS flips"),
            ("💬 Conversational AI Chatbot", "Multi-turn spiritual dialogue with session memory & offline rule fallback"),
            ("📊 4 Role-Based Dashboards", "Tailored consoles for Seeker, Tarot Reader, Spiritual Consultant, Super Admin"),
            ("📄 Multi-Format Dossier Export", "One-click export of complete spiritual guidance dossiers to PDF, DOCX, CSV & JSON"),
            ("🧪 100% Verified Quality", "63/63 pytest automated tests passing in 3.11s + Docker Compose deploy")
        ])
    ]

    for m_title, m_col, mx, my, m_points in milestones:
        # Card background
        m_card = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(mx), Inches(my), Inches(5.75), Inches(2.72))
        m_card.fill.solid()
        m_card.fill.fore_color.rgb = CARD_BG
        m_card.line.color.rgb = CARD_BORDER
        m_card.line.width = Pt(1.2)

        # Title & Badge
        tb_title = slide6.shapes.add_textbox(Inches(mx + 0.15), Inches(my + 0.08), Inches(4.2), Inches(0.32))
        tf_t = tb_title.text_frame
        tf_t.word_wrap = True
        tf_t.margin_left = tf_t.margin_top = 0
        p_t = tf_t.paragraphs[0]
        p_t.text = m_title
        p_t.font.size = Pt(11)
        p_t.font.bold = True
        p_t.font.color.rgb = m_col

        # Status Badge
        b_box = slide6.shapes.add_textbox(Inches(mx + 4.2), Inches(my + 0.08), Inches(1.4), Inches(0.32))
        tf_b = b_box.text_frame
        tf_b.margin_right = tf_b.margin_top = 0
        p_b = tf_b.paragraphs[0]
        p_b.text = "● 100% DELIVERED"
        p_b.alignment = PP_ALIGN.RIGHT
        p_b.font.size = Pt(8.5)
        p_b.font.bold = True
        p_b.font.color.rgb = EMERALD

        # Divider line inside card
        div = slide6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(mx + 0.15), Inches(my + 0.42), Inches(5.45), Inches(0.015))
        div.fill.solid()
        div.fill.fore_color.rgb = CARD_BORDER
        div.line.fill.background()

        # Bullets
        tb_body = slide6.shapes.add_textbox(Inches(mx + 0.15), Inches(my + 0.48), Inches(5.45), Inches(2.15))
        tf_body = tb_body.text_frame
        tf_body.word_wrap = True
        tf_body.margin_left = tf_body.margin_top = 0

        for i, (b_lead, b_desc) in enumerate(m_points):
            p_pt = tf_body.paragraphs[0] if i == 0 else tf_body.add_paragraph()
            p_pt.text = f"▪ {b_lead}: "
            p_pt.font.size = Pt(9.2)
            p_pt.font.bold = True
            p_pt.font.color.rgb = TEXT_WHITE
            run = p_pt.add_run()
            run.text = b_desc
            run.font.bold = False
            run.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 7: Quality Engineering, Test Verification & Future Roadmap
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    apply_bg(slide7)
    add_header(slide7, "Quality Engineering, Test Verification & Future Roadmap", "VERIFICATION & IMPACT", CYAN, slide_num=7)

    # Left Column: Quality & Test Engineering (100% Pass)
    c_qual = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.75), Inches(5.6))
    c_qual.fill.solid()
    c_qual.fill.fore_color.rgb = CARD_BG
    c_qual.line.color.rgb = CARD_BORDER
    c_qual.line.width = Pt(1.2)

    # Title & Badge
    tb_q = slide7.shapes.add_textbox(Inches(1.0), Inches(1.65), Inches(3.8), Inches(0.35))
    tf_q = tb_q.text_frame
    tf_q.margin_left = tf_q.margin_top = 0
    p = tf_q.paragraphs[0]
    p.text = "🧪 Automated Test Suite & Verification"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = CYAN

    badge_q = slide7.shapes.add_textbox(Inches(4.6), Inches(1.65), Inches(1.8), Inches(0.35))
    tf_bq = badge_q.text_frame
    tf_bq.margin_right = tf_bq.margin_top = 0
    pbq = tf_bq.paragraphs[0]
    pbq.text = "● 63/63 PASS (100%)"
    pbq.alignment = PP_ALIGN.RIGHT
    pbq.font.size = Pt(9.5)
    pbq.font.bold = True
    pbq.font.color.rgb = EMERALD

    # Divider
    div_q = slide7.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(2.05), Inches(5.35), Inches(0.015))
    div_q.fill.solid()
    div_q.fill.fore_color.rgb = CARD_BORDER
    div_q.line.fill.background()

    # Content Box
    tb_qc = slide7.shapes.add_textbox(Inches(1.0), Inches(2.15), Inches(5.35), Inches(4.8))
    tf_qc = tb_qc.text_frame
    tf_qc.word_wrap = True
    tf_qc.margin_left = tf_qc.margin_top = 0

    test_points = [
        ("Pytest Execution Speed", "All 63 unit and integration tests execute cleanly in ~3.11s with 0 failures"),
        ("Auth & Security Verification", "Validates bcrypt hashing, JWT token expiry, blacklisting & RBAC permission bounds"),
        ("Reading Engine Integrity", "Verifies all 6 tarot spreads, card uniqueness, elemental balance & palm determinism"),
        ("AI Scoring Math Validation", "Ensures strict compliance with 5-factor weighted scoring formula & letter grades (A+ to D)"),
        ("Data Export Fidelity", "Validates complete data export structures for CSV, JSON, and PDF dossiers"),
        ("Resilience & Offline Fallbacks", "Guarantees graceful degradation when Gemini API key or external DBs are offline"),
        ("Strict Schema Enforcement", "Zero runtime data corruption via strict Pydantic v2 request/response validation")
    ]

    for i, (head, detail) in enumerate(test_points):
        p_t = tf_qc.paragraphs[0] if i == 0 else tf_qc.add_paragraph()
        p_t.text = f"▪ {head}: "
        p_t.font.size = Pt(9.5)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_WHITE
        run = p_t.add_run()
        run.text = detail
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED

    # Right Column: Production Infrastructure & Future Scope
    c_road = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.78), Inches(1.5), Inches(5.75), Inches(5.6))
    c_road.fill.solid()
    c_road.fill.fore_color.rgb = CARD_BG
    c_road.line.color.rgb = CARD_BORDER
    c_road.line.width = Pt(1.2)

    tb_r = slide7.shapes.add_textbox(Inches(6.98), Inches(1.65), Inches(3.8), Inches(0.35))
    tf_r = tb_r.text_frame
    tf_r.margin_left = tf_r.margin_top = 0
    p = tf_r.paragraphs[0]
    p.text = "🚀 Production Infra & Future Horizon"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = PURPLE

    badge_r = slide7.shapes.add_textbox(Inches(10.58), Inches(1.65), Inches(1.8), Inches(0.35))
    tf_br = badge_r.text_frame
    tf_br.margin_right = tf_br.margin_top = 0
    pbr = tf_br.paragraphs[0]
    pbr.text = "● ENTERPRISE READY"
    pbr.alignment = PP_ALIGN.RIGHT
    pbr.font.size = Pt(9.5)
    pbr.font.bold = True
    pbr.font.color.rgb = CYAN

    div_r = slide7.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.98), Inches(2.05), Inches(5.35), Inches(0.015))
    div_r.fill.solid()
    div_r.fill.fore_color.rgb = CARD_BORDER
    div_r.line.fill.background()

    tb_rc = slide7.shapes.add_textbox(Inches(6.98), Inches(2.15), Inches(5.35), Inches(4.8))
    tf_rc = tb_rc.text_frame
    tf_rc.word_wrap = True
    tf_rc.margin_left = tf_rc.margin_top = 0

    future_points = [
        ("Docker Container Network", "Single-command containerization for PostgreSQL, MongoDB, Redis & FastAPI services"),
        ("Live WebRTC Palm Tracking", "Drop-in MediaPipe edge machine learning for real-time video landmark analysis"),
        ("Cross-Platform Mobile Apps", "React Native iOS & Android companion applications with haptic vibration card flips"),
        ("Multilingual Neural Voice", "Real-time voice-guided tarot reading & audio affirmations across 12+ languages"),
        ("WebXR Immersive Sanctum", "3D VR spiritual sanctum supporting interactive virtual reality deck interactions"),
        ("Practitioner Marketplace", "Live video consultation booking system with verified professional tarot readers"),
        ("Continuous Cloud Delivery", "Automated GitHub Actions CI/CD pipeline targeting Vercel & Railway/Render")
    ]

    for i, (head, detail) in enumerate(future_points):
        p_r = tf_rc.paragraphs[0] if i == 0 else tf_rc.add_paragraph()
        p_r.text = f"▪ {head}: "
        p_r.font.size = Pt(9.5)
        p_r.font.bold = True
        p_r.font.color.rgb = TEXT_WHITE
        run = p_r.add_run()
        run.text = detail
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED

    # Save outputs across all target presentation files
    target_files = [
        os.path.join(base_dir, "MysticAI_Presentation.pptx"),
        os.path.join(base_dir, "MysticAI_Presentation [Autosaved].pptx"),
        os.path.join(base_dir, "MysticAI_CrazyAss_Presentation.pptx"),
        os.path.join(base_dir, "docs", "MysticAI_Presentation.pptx"),
        os.path.join(base_dir, "docs", "MysticAI_CrazyAss_Presentation.pptx")
    ]

    for path in target_files:
        try:
            prs.save(path)
            print(f"Saved successfully: {path}")
        except Exception as e:
            print(f"Warning: Could not save to {path}: {e}")

    return target_files[0]

if __name__ == "__main__":
    build_presentation()

